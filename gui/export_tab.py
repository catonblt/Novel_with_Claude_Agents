"""
Export Tab - Export to .docx and compare versions
"""

import customtkinter as ctk
from tkinter import messagebox, filedialog
from pathlib import Path
from typing import Optional, List, Dict
import difflib

try:
    from docx import Document
    from docx.shared import Pt, Inches
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

from .project_manager import ProjectManager
from .utils import format_timestamp


class ExportTab:
    """Export tab for document export and version comparison"""

    def __init__(self, parent, project_manager: ProjectManager):
        self.parent = parent
        self.project_manager = project_manager

        self.versions: List[Dict] = []
        self.selected_versions: List[str] = []

        self._create_widgets()

    def _create_widgets(self):
        """Create tab widgets"""
        # Main container with left and right panels
        main_container = ctk.CTkFrame(self.parent)
        main_container.pack(fill="both", expand=True, padx=10, pady=10)

        # Left panel - Export
        left_panel = ctk.CTkFrame(main_container)
        left_panel.pack(side="left", fill="both", expand=True, padx=(0, 5))

        # Right panel - Versions
        right_panel = ctk.CTkFrame(main_container)
        right_panel.pack(side="right", fill="both", expand=True, padx=(5, 0))

        # === LEFT PANEL: Export ===

        ctk.CTkLabel(
            left_panel,
            text="Export Story",
            font=("Arial", 16, "bold")
        ).pack(anchor="w", padx=10, pady=10)

        # Export section
        export_frame = ctk.CTkFrame(left_panel)
        export_frame.pack(fill="x", padx=10, pady=10)

        ctk.CTkLabel(
            export_frame,
            text="Export final story to Word document (.docx)",
            font=("Arial", 11)
        ).pack(anchor="w", padx=10, pady=10)

        # Export options
        options_frame = ctk.CTkFrame(export_frame)
        options_frame.pack(fill="x", padx=10, pady=10)

        ctk.CTkLabel(options_frame, text="Document Title:").pack(anchor="w", pady=(0, 5))
        self.doc_title_entry = ctk.CTkEntry(
            options_frame,
            placeholder_text="Will use project name if blank"
        )
        self.doc_title_entry.pack(fill="x", pady=(0, 10))

        ctk.CTkLabel(options_frame, text="Author Name:").pack(anchor="w", pady=(0, 5))
        self.doc_author_entry = ctk.CTkEntry(
            options_frame,
            placeholder_text="Will use config author if blank"
        )
        self.doc_author_entry.pack(fill="x", pady=(0, 10))

        # Include agent outputs checkbox
        self.include_agents_var = ctk.BooleanVar(value=False)
        include_agents_cb = ctk.CTkCheckBox(
            options_frame,
            text="Include all agent outputs as appendix",
            variable=self.include_agents_var
        )
        include_agents_cb.pack(anchor="w", pady=10)

        # Export button
        export_btn = ctk.CTkButton(
            export_frame,
            text="Export to Word Document",
            command=self._export_to_docx,
            height=40,
            font=("Arial", 12, "bold")
        )
        export_btn.pack(fill="x", padx=10, pady=10)

        if not DOCX_AVAILABLE:
            export_btn.configure(state="disabled")
            ctk.CTkLabel(
                export_frame,
                text="⚠ python-docx not installed. Run: pip install python-docx",
                text_color="orange",
                font=("Arial", 10)
            ).pack(padx=10, pady=5)

        # === RIGHT PANEL: Versions ===

        ctk.CTkLabel(
            right_panel,
            text="Version Management",
            font=("Arial", 16, "bold")
        ).pack(anchor="w", padx=10, pady=10)

        # Create version section
        version_create_frame = ctk.CTkFrame(right_panel)
        version_create_frame.pack(fill="x", padx=10, pady=10)

        ctk.CTkLabel(
            version_create_frame,
            text="Create New Version Snapshot:",
            font=("Arial", 11)
        ).pack(anchor="w", padx=10, pady=(10, 5))

        version_name_frame = ctk.CTkFrame(version_create_frame)
        version_name_frame.pack(fill="x", padx=10, pady=5)

        self.version_name_entry = ctk.CTkEntry(
            version_name_frame,
            placeholder_text="Optional version name"
        )
        self.version_name_entry.pack(side="left", fill="x", expand=True, padx=(0, 10))

        create_version_btn = ctk.CTkButton(
            version_name_frame,
            text="Create",
            command=self._create_version,
            width=80
        )
        create_version_btn.pack(side="right")

        # Version list
        ctk.CTkLabel(
            right_panel,
            text="Saved Versions:",
            font=("Arial", 12, "bold")
        ).pack(anchor="w", padx=10, pady=(20, 5))

        # Refresh button
        refresh_btn = ctk.CTkButton(
            right_panel,
            text="Refresh List",
            command=self._refresh_versions,
            width=100
        )
        refresh_btn.pack(anchor="e", padx=10, pady=(0, 5))

        # Versions list
        self.versions_frame = ctk.CTkScrollableFrame(right_panel)
        self.versions_frame.pack(fill="both", expand=True, padx=10, pady=(0, 10))

        # Compare button
        self.compare_btn = ctk.CTkButton(
            right_panel,
            text="Compare Selected Versions",
            command=self._compare_versions,
            state="disabled"
        )
        self.compare_btn.pack(fill="x", padx=10, pady=(0, 10))

    def _export_to_docx(self):
        """Export story to Word document"""
        if not DOCX_AVAILABLE:
            messagebox.showerror(
                "Error",
                "python-docx library not installed. Install with: pip install python-docx"
            )
            return

        if not self.project_manager.current_project_dir:
            messagebox.showwarning("Warning", "No project loaded")
            return

        # Load final story
        success, content, message = self.project_manager.load_final_story()
        if not success or not content:
            messagebox.showwarning("Warning", "No final story to export. Generate content first.")
            return

        # Get export location
        default_name = self.project_manager.get_config().get("project_name", "novel").replace(" ", "_")
        file_path = filedialog.asksaveasfilename(
            defaultextension=".docx",
            filetypes=[("Word Document", "*.docx")],
            initialfile=f"{default_name}.docx"
        )

        if not file_path:
            return

        try:
            # Create document
            doc = Document()

            # Get metadata
            config = self.project_manager.get_config()
            title = self.doc_title_entry.get().strip() or config.get("project_name", "Untitled Novel")
            author = self.doc_author_entry.get().strip() or config.get("metadata", {}).get("author", "Unknown Author")

            # Set document properties
            doc.core_properties.title = title
            doc.core_properties.author = author

            # Add title
            title_para = doc.add_heading(title, level=0)
            title_para.alignment = 1  # Center

            # Add author
            author_para = doc.add_paragraph(f"by {author}")
            author_para.alignment = 1  # Center

            doc.add_page_break()

            # Add story content
            # Parse markdown-style content
            for line in content.split('\n'):
                line = line.rstrip()

                if line.startswith('# '):
                    doc.add_heading(line[2:], level=1)
                elif line.startswith('## '):
                    doc.add_heading(line[3:], level=2)
                elif line.startswith('### '):
                    doc.add_heading(line[4:], level=3)
                elif line.strip() == '':
                    doc.add_paragraph()
                else:
                    doc.add_paragraph(line)

            # Add agent outputs if requested
            if self.include_agents_var.get():
                doc.add_page_break()
                doc.add_heading("Appendix: Agent Outputs", level=1)

                for agent_num in sorted(self.project_manager.AGENTS.keys()):
                    success, agent_content, _ = self.project_manager.load_agent_output(agent_num)
                    if success and agent_content:
                        agent_name = self.project_manager.AGENTS[agent_num]["name"]
                        doc.add_heading(f"Agent {agent_num}: {agent_name}", level=2)

                        for line in agent_content.split('\n'):
                            if line.strip():
                                doc.add_paragraph(line)

                        doc.add_page_break()

            # Save document
            doc.save(file_path)

            messagebox.showinfo("Success", f"Document exported to:\n{file_path}")

        except Exception as e:
            messagebox.showerror("Error", f"Failed to export document:\n{str(e)}")

    def _create_version(self):
        """Create a new version snapshot"""
        if not self.project_manager.current_project_dir:
            messagebox.showwarning("Warning", "No project loaded")
            return

        version_name = self.version_name_entry.get().strip()
        success, message = self.project_manager.create_version(version_name or None)

        if success:
            messagebox.showinfo("Success", message)
            self.version_name_entry.delete(0, "end")
            self._refresh_versions()
        else:
            messagebox.showerror("Error", message)

    def _refresh_versions(self):
        """Refresh the versions list"""
        # Clear current list
        for widget in self.versions_frame.winfo_children():
            widget.destroy()

        self.selected_versions = []

        # Get versions
        self.versions = self.project_manager.list_versions()

        if not self.versions:
            ctk.CTkLabel(
                self.versions_frame,
                text="No versions created yet",
                text_color="gray"
            ).pack(pady=20)
            return

        # Display versions
        for version in self.versions:
            self._add_version_item(version)

    def _add_version_item(self, version: Dict):
        """Add a version item to the list"""
        item_frame = ctk.CTkFrame(self.versions_frame)
        item_frame.pack(fill="x", pady=5)

        # Checkbox for selection
        var = ctk.BooleanVar()
        cb = ctk.CTkCheckBox(
            item_frame,
            text="",
            variable=var,
            command=lambda: self._on_version_select(version['version_name'], var.get()),
            width=20
        )
        cb.pack(side="left", padx=5)

        # Version info
        info_frame = ctk.CTkFrame(item_frame)
        info_frame.pack(side="left", fill="x", expand=True, padx=5)

        name_label = ctk.CTkLabel(
            info_frame,
            text=version['version_name'],
            font=("Arial", 11, "bold"),
            anchor="w"
        )
        name_label.pack(anchor="w")

        date_label = ctk.CTkLabel(
            info_frame,
            text=f"Created: {version['created']}",
            font=("Arial", 9),
            text_color="gray",
            anchor="w"
        )
        date_label.pack(anchor="w")

    def _on_version_select(self, version_name: str, selected: bool):
        """Handle version selection"""
        if selected:
            if version_name not in self.selected_versions:
                self.selected_versions.append(version_name)
        else:
            if version_name in self.selected_versions:
                self.selected_versions.remove(version_name)

        # Enable compare button if exactly 2 versions selected
        if len(self.selected_versions) == 2:
            self.compare_btn.configure(state="normal")
        else:
            self.compare_btn.configure(state="disabled")

    def _compare_versions(self):
        """Compare two selected versions"""
        if len(self.selected_versions) != 2:
            messagebox.showwarning("Warning", "Please select exactly 2 versions to compare")
            return

        # Get version data
        version1_name = self.selected_versions[0]
        version2_name = self.selected_versions[1]

        version1 = next((v for v in self.versions if v['version_name'] == version1_name), None)
        version2 = next((v for v in self.versions if v['version_name'] == version2_name), None)

        if not version1 or not version2:
            messagebox.showerror("Error", "Could not load version data")
            return

        # Load story content from both versions
        try:
            v1_story_path = Path(version1['path']) / "manuscript" / "final_story.md"
            v2_story_path = Path(version2['path']) / "manuscript" / "final_story.md"

            v1_content = v1_story_path.read_text() if v1_story_path.exists() else "[No content]"
            v2_content = v2_story_path.read_text() if v2_story_path.exists() else "[No content]"

            # Show comparison window
            ComparisonWindow(self.parent, version1_name, v1_content, version2_name, v2_content)

        except Exception as e:
            messagebox.showerror("Error", f"Failed to compare versions:\n{str(e)}")

    def refresh(self):
        """Refresh the export tab"""
        self._refresh_versions()

        # Update default values from config
        config = self.project_manager.get_config()
        if config:
            self.doc_author_entry.delete(0, "end")
            self.doc_author_entry.insert(0, config.get("metadata", {}).get("author", ""))


class ComparisonWindow(ctk.CTkToplevel):
    """Window for comparing two versions side-by-side"""

    def __init__(self, parent, version1_name: str, content1: str, version2_name: str, content2: str):
        super().__init__(parent)

        self.title("Compare Versions")
        self.geometry("1200x700")

        # Title
        title = ctk.CTkLabel(
            self,
            text=f"Comparing: {version1_name} vs {version2_name}",
            font=("Arial", 16, "bold")
        )
        title.pack(pady=10)

        # Main comparison frame
        compare_frame = ctk.CTkFrame(self)
        compare_frame.pack(fill="both", expand=True, padx=10, pady=10)

        # Left panel - Version 1
        left_panel = ctk.CTkFrame(compare_frame)
        left_panel.pack(side="left", fill="both", expand=True, padx=(0, 5))

        ctk.CTkLabel(
            left_panel,
            text=version1_name,
            font=("Arial", 12, "bold")
        ).pack(pady=5)

        text1 = ctk.CTkTextbox(left_panel, font=("Courier", 10))
        text1.pack(fill="both", expand=True)
        text1.insert("1.0", content1)
        text1.configure(state="disabled")

        # Right panel - Version 2
        right_panel = ctk.CTkFrame(compare_frame)
        right_panel.pack(side="right", fill="both", expand=True, padx=(5, 0))

        ctk.CTkLabel(
            right_panel,
            text=version2_name,
            font=("Arial", 12, "bold")
        ).pack(pady=5)

        text2 = ctk.CTkTextbox(right_panel, font=("Courier", 10))
        text2.pack(fill="both", expand=True)
        text2.insert("1.0", content2)
        text2.configure(state="disabled")

        # Close button
        close_btn = ctk.CTkButton(self, text="Close", command=self.destroy, width=100)
        close_btn.pack(pady=10)
